"""
populate_malayalam_part3.py

Parses assets/db/Part III_word.docx and inserts Malayalam Quran data
for surahs 19-36 into assets/db/quran_asad_combined_nw.sqlite.

Populates tables: malayalam_surahs, malayalam_verses
(No footnotes in Part III document)

Run from repo root:
    python scripts/populate_malayalam_part3.py
"""

import os
import sys
import re
import shutil
import sqlite3

from docx import Document

# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
REPO_ROOT = os.path.dirname(SCRIPT_DIR)
DB_PATH = os.path.join(REPO_ROOT, "assets", "db", "quran_asad_combined_nw.sqlite")
DOCX_PATH = os.path.join(REPO_ROOT, "assets", "db", "Part III_word.docx")
BACKUP_PATH = DB_PATH + ".bak_part3"

# Sidecar files that must be removed to avoid corruption
WAL_PATH = DB_PATH + "-wal"
SHM_PATH = DB_PATH + "-shm"

# ---------------------------------------------------------------------------
# Surah metadata for surahs 19-36
# (arabic_name, english_translation)
# ---------------------------------------------------------------------------
SURAH_META = {
    19: ("Maryam",       "Mary"),
    20: ("Ta-Ha",        "O Man"),
    21: ("Al-Anbiya",    "The Prophets"),
    22: ("Al-Hajj",      "The Pilgrimage"),
    23: ("Al-Mu'minun",  "The Believers"),
    24: ("An-Nur",       "The Light"),
    25: ("Al-Furqan",    "The Criterion"),
    26: ("Ash-Shu'ara",  "The Poets"),
    27: ("An-Naml",      "The Ants"),
    28: ("Al-Qasas",     "The Story"),
    29: ("Al-Ankabut",   "The Spider"),
    30: ("Ar-Rum",       "The Romans"),
    31: ("Luqman",       "Luqman"),
    32: ("As-Sajda",     "The Prostration"),
    33: ("Al-Ahzab",     "The Confederates"),
    34: ("Saba",         "Sheba"),
    35: ("Al-Fatir",     "The Creator"),
    36: ("Ya-Sin",       "Ya Sin"),
}

# Verse-number prefix pattern.
# Handles all observed variants in the document:
#   Normal:    "19.1 "   "19:1 "   "(19.6) "  "(19:11) "
#   Space:     "26: 20 "
#   Semicolon: "26;72 "  "28;5 "
#   Hyphen:    "26-59 "
#   Bracket:   "[27:38 "
VERSE_PREFIX_RE = re.compile(r"^\[?\(?(\d+)[-.:;]\s*(\d+)\)?\]?\s*")

# Hardcoded verse-number corrections for specific document typos.
# Key: (current_chapter, raw_surah, raw_verse) → correct_verse_number
# Only needed when the sequential-fallback heuristic cannot infer the right number.
VERSE_TYPO_FIXES: dict = {
    # "27:8" appears in chapter 28's section between v6 and v8 — should be v7.
    # raw_verse=8 is within +4 of last_verse=6, so it would be accepted as 8 without this fix.
    (28, 27, 8): 7,
}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def get_font_size(para):
    """Return font size (pt) of the paragraph's first non-empty run, or
    fall back to the paragraph style's font size.  Returns None if unknown."""
    for run in para.runs:
        if run.text.strip() and run.font.size:
            return round(run.font.size.pt, 1)
    if para.style and para.style.font.size:
        return round(para.style.font.size.pt, 1)
    return None


def is_normal(para):
    return para.style and para.style.name == "Normal"


# ---------------------------------------------------------------------------
# Parse docx
# ---------------------------------------------------------------------------

def parse_docx(path):
    """Return a list of surah dicts with keys:
        chapter_number, arabic_name, malayalam_name, english_translation,
        revelation_period, introduction, verses (list of (verse_num, text))
    """
    doc = Document(path)
    paras = doc.paragraphs

    surahs = []
    current = None
    surah_number = 18  # will be incremented to 19 on first heading
    collecting_intro = False

    # Per-surah verse tracking (reset on each new surah heading)
    last_verse_num = 0
    seen_verse_nums: set = set()

    for para in paras:
        text = para.text.strip()
        if not text:
            continue

        sz = get_font_size(para)
        normal = is_normal(para)

        # ---- Surah heading: Normal, ~18 pt ---------------------------------
        if normal and sz is not None and abs(sz - 18.0) < 0.6:
            surah_number += 1
            if surah_number > 36:
                break  # done with Part III

            arabic_name, english_trans = SURAH_META.get(
                surah_number, ("Unknown", "Unknown")
            )
            current = {
                "chapter_number": surah_number,
                "arabic_name": arabic_name,
                "malayalam_name": text,
                "english_translation": english_trans,
                "revelation_period": "",
                "intro_parts": [],  # collect 9pt paragraphs
                "verses": [],
            }
            surahs.append(current)
            collecting_intro = False
            last_verse_num = 0
            seen_verse_nums = set()
            continue

        if current is None:
            continue

        # ---- Revelation period: Normal, ~8 pt ------------------------------
        if normal and sz is not None and abs(sz - 8.0) < 0.6:
            if not current["revelation_period"]:
                current["revelation_period"] = text
                collecting_intro = True
            continue

        # ---- Introduction text: Normal, ~9 pt ------------------------------
        if normal and sz is not None and abs(sz - 9.0) < 0.6 and collecting_intro:
            current["intro_parts"].append(text)
            continue

        # ---- Arabic verse / Bismillah: Normal, ~16 pt ----------------------
        if normal and sz is not None and abs(sz - 16.0) < 0.6:
            collecting_intro = False  # intro ends when Arabic text begins
            continue

        # ---- Malayalam verse translation: Normal, ~9.5 pt (or 12 pt) ------
        if normal and sz is not None and (
            abs(sz - 9.5) < 0.6 or abs(sz - 12.0) < 0.6
        ):
            m = VERSE_PREFIX_RE.match(text)
            if not m:
                # No verse-number prefix → Bismillah translation or similar; skip
                continue

            raw_surah = int(m.group(1))
            raw_verse = int(m.group(2))
            verse_text = text[m.end():]

            # Check for combined-verse range suffix (e.g., "34:30-31 text").
            # If present, we will store the text under EVERY verse in the range.
            combined_end: int | None = None
            range_m = re.match(r"^-(\d+)\s*", verse_text)
            if range_m:
                combined_end = int(range_m.group(1))
                verse_text = verse_text[range_m.end():]

            # Accept only paragraphs whose surah tag is within ±15 of the current
            # chapter (catches common surah-number typos in the document).
            if abs(raw_surah - current["chapter_number"]) > 15:
                continue

            # Apply hardcoded verse-number corrections
            fix_key = (current["chapter_number"], raw_surah, raw_verse)
            if fix_key in VERSE_TYPO_FIXES:
                raw_verse = VERSE_TYPO_FIXES[fix_key]

            # Decide actual verse number:
            # • If raw_verse is the natural next verse (1–4 ahead of last),
            #   trust it as-is.
            # • Otherwise (large jump forward, backwards, or duplicate) fall
            #   back to sequential: last_verse_num + 1.
            if (
                raw_verse > last_verse_num
                and raw_verse <= last_verse_num + 4
                and raw_verse not in seen_verse_nums
            ):
                actual_verse = raw_verse
            else:
                # Sequential fallback — skip any already-seen numbers
                actual_verse = last_verse_num + 1
                while actual_verse in seen_verse_nums:
                    actual_verse += 1

            current["verses"].append((actual_verse, verse_text))
            seen_verse_nums.add(actual_verse)
            last_verse_num = actual_verse

            # Also store under every additional verse in a combined range (e.g. 30-31)
            if combined_end and combined_end > actual_verse:
                for extra_verse in range(actual_verse + 1, combined_end + 1):
                    if extra_verse not in seen_verse_nums:
                        current["verses"].append((extra_verse, verse_text))
                        seen_verse_nums.add(extra_verse)
                        last_verse_num = extra_verse
            continue

    # Build final introduction string for each surah
    for surah in surahs:
        parts = surah.pop("intro_parts")
        rev = surah["revelation_period"]
        if parts:
            surah["introduction"] = rev + "\n\n" + "\n".join(parts)
        else:
            surah["introduction"] = rev

    return surahs


# ---------------------------------------------------------------------------
# Database insertion
# ---------------------------------------------------------------------------

def insert_into_db(db_path, surahs):
    con = sqlite3.connect(db_path)
    c = con.cursor()

    # Use DELETE journal mode (required by app)
    c.execute("PRAGMA journal_mode=DELETE")

    # Remove existing data for surahs 19-36 (idempotency)
    print("Removing any existing rows for chapter_number 19-36 ...")

    # Find surah IDs for chapters 19-36
    c.execute(
        "SELECT id FROM malayalam_surahs WHERE chapter_number >= 19 AND chapter_number <= 36"
    )
    existing_surah_ids = [row[0] for row in c.fetchall()]
    if existing_surah_ids:
        placeholders = ",".join("?" for _ in existing_surah_ids)
        c.execute(
            f"DELETE FROM malayalam_verses WHERE surah_id IN ({placeholders})",
            existing_surah_ids,
        )
    c.execute(
        "DELETE FROM malayalam_surahs WHERE chapter_number >= 19 AND chapter_number <= 36"
    )
    con.commit()

    # Insert surahs and verses
    total_verses = 0
    for surah in surahs:
        cn = surah["chapter_number"]
        print(
            f"  Inserting Surah {cn} ({surah['arabic_name']}): "
            f"{len(surah['verses'])} verses ..."
        )

        c.execute(
            """
            INSERT INTO malayalam_surahs
                (id, chapter_number, arabic_name, malayalam_name,
                 english_translation, revelation_period, introduction)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                cn,
                cn,
                surah["arabic_name"],
                surah["malayalam_name"],
                surah["english_translation"],
                surah["revelation_period"],
                surah["introduction"],
            ),
        )

        for verse_num, verse_text in surah["verses"]:
            c.execute(
                """
                INSERT INTO malayalam_verses
                    (surah_id, verse_number, malayalam_translation)
                VALUES (?, ?, ?)
                """,
                (cn, verse_num, verse_text),
            )
        total_verses += len(surah["verses"])

    con.commit()
    con.close()
    return total_verses


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    print(f"DB   : {DB_PATH}")
    print(f"Docx : {DOCX_PATH}")

    if not os.path.exists(DOCX_PATH):
        sys.exit(f"ERROR: Word document not found: {DOCX_PATH}")
    if not os.path.exists(DB_PATH):
        sys.exit(f"ERROR: Database not found: {DB_PATH}")

    # --- Backup ---
    print(f"\nBacking up DB to: {BACKUP_PATH}")
    shutil.copy2(DB_PATH, BACKUP_PATH)

    # --- Remove WAL/SHM sidecar files ---
    for sidecar in (WAL_PATH, SHM_PATH):
        if os.path.exists(sidecar):
            os.remove(sidecar)
            print(f"Removed sidecar: {sidecar}")

    # --- Parse ---
    print("\nParsing Word document ...")
    surahs = parse_docx(DOCX_PATH)

    print(f"Parsed {len(surahs)} surahs:")
    for s in surahs:
        print(
            f"  Surah {s['chapter_number']:2d} {s['arabic_name']:<15} "
            f"| {len(s['verses']):3d} verses "
            f"| rev_period: {s['revelation_period'][:30]!r}"
        )

    if len(surahs) != 18:
        print(
            f"\nWARNING: Expected 18 surahs (19-36), got {len(surahs)}. "
            "Check parsing output above before proceeding."
        )
        resp = input("Continue anyway? [y/N] ")
        if resp.strip().lower() != "y":
            sys.exit("Aborted.")

    # --- Insert ---
    print("\nInserting into DB ...")
    total_verses = insert_into_db(DB_PATH, surahs)

    # --- Verify ---
    print("\n--- Verification ---")
    con = sqlite3.connect(DB_PATH)
    c = con.cursor()

    c.execute("SELECT COUNT(*) FROM malayalam_surahs")
    surah_count = c.fetchone()[0]
    c.execute("SELECT COUNT(*) FROM malayalam_verses")
    verse_count = c.fetchone()[0]
    c.execute("SELECT COUNT(*) FROM malayalam_footnotes")
    fn_count = c.fetchone()[0]

    c.execute(
        "SELECT chapter_number, arabic_name, malayalam_name FROM malayalam_surahs "
        "WHERE chapter_number >= 19 ORDER BY chapter_number"
    )
    new_surahs = c.fetchall()
    con.close()

    print(f"malayalam_surahs  total rows : {surah_count}")
    print(f"malayalam_verses  total rows : {verse_count}")
    print(f"malayalam_footnotes total rows: {fn_count}")
    print(f"\nNewly inserted surahs ({len(new_surahs)}):")
    for row in new_surahs:
        print(f"  ch={row[0]:2d}  arabic={row[1]:<15}  malayalam={row[2][:40]}")

    print(f"\nDone. Inserted {len(surahs)} surahs and {total_verses} verses.")


if __name__ == "__main__":
    main()
